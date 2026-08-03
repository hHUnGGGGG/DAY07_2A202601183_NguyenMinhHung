"""
make_phancong_xlsx.py — sinh PHAN_CONG_NHOM.xlsx (1 sheet) từ bảng phân công dưới đây.

    pip install openpyxl
    py scripts/make_phancong_xlsx.py

Sửa danh sách ROWS rồi chạy lại để cập nhật file. Cột "Trạng thái" có dropdown.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_PATH = Path(__file__).resolve().parent.parent / "PHAN_CONG_NHOM.xlsx"

HEADERS = [
    "STT",
    "Thành viên",
    "Vai điều phối",
    "Chiến lược chunking",
    "Nhiệm vụ",
    "File code / file phụ trách",
    "Deliverable (bằng chứng)",
    "Mốc",
    "Phụ thuộc",
    "Trạng thái",
]

WIDTHS = [5, 20, 18, 30, 52, 40, 34, 8, 26, 15]

ALL = "Cả 5 người"
TV1, TV2, TV3, TV4, TV5 = (
    "Nguyễn Minh Hùng",
    "Hoàng Anh Quân",
    "Phạm Hải Đăng",
    "Bùi Gia Huy",
    "Nguyễn Huy Đức",
)

ROLES = {
    ALL: "—",
    TV1: "Repo owner / Đồng bộ",
    TV2: "Benchmark owner",
    TV3: "Data curator",
    TV4: "QA / Integration",
    TV5: "Demo coordinator",
}

STRATEGIES = {
    ALL: "—",
    TV1: "FixedSizeChunker(500, 50)",
    TV2: "SentenceChunker(3 câu)",
    TV3: "RecursiveChunker(400)",
    TV4: "HeadingChunker(800)",
    TV5: "HeadingRecursiveChunker(800, breadcrumb)",
}

# (thành viên, nhiệm vụ, file, deliverable, mốc, phụ thuộc)
ROWS = [
    # ---------------------------------------------------------------- việc chung
    (ALL, "Tạo repo riêng DAY07-<MSSV>-<HoTen>, tạo venv, pip install -r requirements.txt",
     "(môi trường)", "repo chạy được", "CP0", "—"),
    (ALL, "Chạy self-check pipeline nạp dữ liệu có sẵn",
     "ingest.py", "output 'ingest self-check OK'", "CP0", "—"),
    (ALL, "Sync 3 artifact DÙNG CHUNG từ repo nguồn của Hùng; không sửa cục bộ",
     "data/k3_university/, bench_queries.py, report/REPORT_NHOM.md", "3 artifact giống repo nguồn",
     "CP0", "TV1 đã push"),
    (ALL, "Dịch phần tài liệu được giao sang tiếng Việt (giữ nguyên số + thuật ngữ gốc trong ngoặc)",
     "data/k3_university/*.md", "file .md hết marker TODO", "CP2", "TV3 chốt nguồn"),
    (ALL, "Warm-up: cosine similarity + bài toán chunking (10.000 ký tự → 23 chunk)",
     "report/REPORT_CANHAN.md", "2 câu trả lời, viết TRƯỚC khi code", "CP3", "—"),
    (ALL, "TASK 1-3: SentenceChunker, RecursiveChunker + _split, compute_similarity, Comparator",
     "src/chunking.py", "pytest -k 'Chunker or Similarity or Compare' → 23 passed", "CP3", "—"),
    (ALL, "TASK 4-5: EmbeddingStore (_make_record, _search_records, add, search, size, filter, delete)",
     "src/store.py", "14 test store pass", "CP4", "CP3"),
    (ALL, "TASK 6: KnowledgeBaseAgent — context đánh số [1][2][3] kèm doc_id",
     "src/agent.py", "agent trả về string không rỗng", "CP4", "CP3"),
    (ALL, "Chạy pytest đầy đủ + main.py, dán output THẬT vào report cá nhân (30 điểm)",
     "report/REPORT_CANHAN.md", "42 passed", "CP4", "CP3"),
    (ALL, "Chạy bench.py với strategy riêng của mình (chỉ đổi --strategy)",
     "bench.py → report/bench_<strategy>.md", "5 query có top-3 + agent answer", "CP5", "CP4 + CP2"),
    (ALL, "Dự đoán 5 cặp câu similarity, so với thực tế, viết reflection",
     "report/REPORT_CANHAN.md", "bảng dự đoán vs thực tế", "CP6", "CP3"),
    (ALL, "Phân tích 5 tiêu chí + ít nhất 1 failure case CÓ BẰNG CHỨNG từ top-k",
     "report/REPORT_CANHAN.md", "failure case đủ 4 phần: query→bằng chứng→nguyên nhân→đề xuất",
     "CP6", "CP5"),

    # ---------------------------------------------------------------- TV1 Hùng
    (TV1, "Giữ repo nguồn: push corpus + bench_queries.py, điều phối 4 người sync lại",
     "(git)", "5 repo cùng corpus hash", "CP2", "—"),
    (TV1, "Viết harness benchmark dùng chung + registry 5 strategy",
     "bench.py", "chạy được cả 5 --strategy", "CP5", "bench_queries.py của Quân"),
    (TV1, "Dịch tài liệu #1 Thư viện và #8 Thông báo đăng ký học phần",
     "data/k3_university/vinuni-library-access-services.md, vinuni-course-registration-announcement.md",
     "2 file đủ metadata", "CP2", "—"),
    (TV1, "Chạy ChunkingStrategyComparator trên 2-3 tài liệu (BỎ front matter trước khi so sánh)",
     "report/REPORT_NHOM.md §2", "bảng baseline 3 strategy", "CP5", "CP3"),

    # ---------------------------------------------------------------- TV2 Quân
    (TV2, "Soạn 5 benchmark query: số liệu / điều kiện / quy trình / liệt kê / ngoại lệ+filter",
     "bench_queries.py", "validate() không còn vấn đề", "CP5", "CP2"),
    (TV2, "Trích gold answer + must_contain từ corpus, đối chiếu văn bản gốc, đặt verified=True",
     "bench_queries.py", "5 query verified", "CP5", "CP2"),
    (TV2, "Dịch tài liệu #2 Quy chế đào tạo (gánh 2/5 query — ưu tiên cao nhất)",
     "data/k3_university/vinuni-academic-regulations-undergrad.md", "file đủ metadata", "CP2", "—"),
    (TV2, "Điền bảng 5 query + gold answer vào report nhóm; KHÔNG đổi query sau khi ai đó đã chạy bench",
     "report/REPORT_NHOM.md §3", "bảng 5 query", "CP5", "—"),

    # ---------------------------------------------------------------- TV3 Đăng
    (TV3, "Chốt phạm vi corpus, kiểm robots.txt/điều khoản, thu thập 8 tài liệu VinUni công khai",
     "docs/CORPUS_PLAN.md", "8 file .md có nội dung thật", "CP2", "—"),
    (TV3, "Giữ sources.csv khớp 1-1 với corpus (doc_id, URL, retrieved_at, document_version)",
     "data/k3_university/sources.csv", "check_corpus.py báo 'khớp 1-1'", "CP2", "—"),
    (TV3, "Dịch tài liệu #3 Quy tắc ứng xử SINH VIÊN — làm CÙNG LÚC với Huy (#6) để giữ khác biệt thật",
     "data/k3_university/vinuni-student-code-of-conduct.md", "mức kỷ luật khác rõ với bản nhân viên",
     "CP2", "—"),
    (TV3, "Điền Data Inventory + Metadata Schema vào report nhóm ngay sau khi CP2 đạt",
     "report/REPORT_NHOM.md §1", "2 bảng đã điền", "CP2", "CP2"),

    # ---------------------------------------------------------------- TV4 Huy
    (TV4, "Viết script tự động hoá CHECKPOINT 2 + in corpus hash",
     "scripts/check_corpus.py", "mọi dòng OK", "CP2", "—"),
    (TV4, "Viết split_into_sections + HeadingChunker (KHÔNG gắn lại tiêu đề — chủ đích)",
     "src/heading_chunker.py", "py -m src.heading_chunker self-check OK", "CP5", "CP3"),
    (TV4, "Dịch tài liệu #6 Quy tắc ứng xử NHÂN VIÊN — document nhiễu chủ đích của Q5",
     "data/k3_university/vinuni-employees-code-of-conduct.md", "cùng từ vựng, khác đáp án so với #3",
     "CP2", "—"),
    (TV4, "Rà checklist CP7 cho cả nhóm: 42 passed, không commit .env/.venv, tên repo đúng quy ước",
     "(review 5 repo)", "checklist CP7 tick đủ", "CP7", "CP6"),

    # ---------------------------------------------------------------- TV5 Đức
    (TV5, "Viết HeadingRecursiveChunker — GẮN LẠI breadcrumb tiêu đề vào mảnh con",
     "src/heading_chunker.py", "self-check thấy mảnh con mang breadcrumb", "CP5", "CP3"),
    (TV5, "Dịch tài liệu #4 Quy định tài chính, #5 Duy trì học bổng, #7 Bổ nhiệm giảng viên",
     "data/k3_university/vinuni-financial-regulations-student.md, vinuni-scholarship-maintenance.md, vinuni-faculty-appointment-promotion.md",
     "3 file đủ metadata", "CP2", "—"),
    (TV5, "Gom bảng so sánh 5 thành viên + kết luận strategy nào hợp chủ đề này (15 điểm)",
     "report/REPORT_NHOM.md §2", "bảng so sánh + 2-3 câu lý giải", "CP6", "CP5 của cả 5"),
    (TV5, "Dựng kịch bản demo 6-8 phút, mở sẵn terminal chạy được bench.py",
     "report/REPORT_NHOM.md §4", "kịch bản 4 phần theo đúng thời lượng", "CP7", "CP6"),
]

NOTES = [
    "src/ MỖI NGƯỜI TỰ CODE trong repo riêng của mình — đây là 30 điểm cá nhân, không ai code hộ ai.",
    "Dùng chung và phải giống hệt nhau giữa 5 repo: data/k3_university/, bench_queries.py, report/REPORT_NHOM.md.",
    "Biến duy nhất được phép khác nhau khi benchmark là chiến lược chunking. Corpus, 5 query, embedding backend, top_k phải giữ nguyên.",
    "Mốc: CP0 khởi tạo · CP2 corpus · CP3 chunking · CP4 42 test pass · CP5 benchmark chạy được · CP6 phân tích · CP7 nộp bài.",
    "Chi tiết từng bước: docs/superpowers/plans/2026-08-03-lab07-k3-vinuni-plan.md",
]

# ------------------------------------------------------------------------------ style
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
MEMBER_FILLS = {
    ALL: "EEF3FA",
    TV1: "FFFFFF",
    TV2: "F7F7F7",
    TV3: "FFFFFF",
    TV4: "F7F7F7",
    TV5: "FFFFFF",
}


def build() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Phân công"

    sheet["A1"] = "PHÂN CÔNG NHÓM — LAB 07: Embedding & Vector Store (biến thể K3 — Quy định VinUniversity)"
    sheet["A1"].font = Font(bold=True, size=14, color="1F3864")
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEADERS))

    row_index = 2
    for note in NOTES:
        sheet.cell(row=row_index, column=1, value="• " + note).font = Font(size=10, italic=True, color="444444")
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=len(HEADERS))
        row_index += 1

    header_row = row_index + 1
    for column_index, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column_index, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    for offset, (member, task, files, deliverable, milestone, depends) in enumerate(ROWS):
        row = header_row + 1 + offset
        values = [
            offset + 1,
            member,
            ROLES[member],
            STRATEGIES[member],
            task,
            files,
            deliverable,
            milestone,
            depends,
            "Chưa bắt đầu",
        ]
        fill = PatternFill("solid", fgColor=MEMBER_FILLS[member])
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column_index, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BORDER
            cell.fill = fill
            if column_index in (2, 3):
                cell.font = Font(bold=(member != ALL))
            if column_index in (1, 8, 10):
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    last_row = header_row + len(ROWS)

    for column_index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(HEADERS))}{last_row}"

    validation = DataValidation(
        type="list",
        formula1='"Chưa bắt đầu,Đang làm,Chờ review,Xong,Bị chặn"',
        allow_blank=True,
        showDropDown=False,
    )
    sheet.add_data_validation(validation)
    validation.add(f"J{header_row + 1}:J{last_row}")

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"{header_row}:{header_row}"

    workbook.save(OUT_PATH)
    print(f"Đã ghi {OUT_PATH} — 1 sheet, {len(ROWS)} dòng công việc.")


if __name__ == "__main__":
    build()
